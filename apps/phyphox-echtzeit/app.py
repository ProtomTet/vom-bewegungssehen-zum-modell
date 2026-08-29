from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from threading import Lock
from typing import Any
from urllib.parse import urlparse

import math
import re
import tempfile

import requests
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font

APP_DIR = Path(__file__).resolve().parent

app = Flask(__name__)


@dataclass
class StreamState:
    base_url: str | None = None
    experiment_title: str | None = None
    session_id: str | None = None
    connected: bool = False
    buffer_options: list[str] = field(default_factory=list)
    buffers: dict[str, str | None] = field(
        default_factory=lambda: {"time": None, "x": None, "y": None, "z": None}
    )
    data: dict[str, list[float | None]] = field(
        default_factory=lambda: {"time": [], "x": [], "y": [], "z": [], "mag": []}
    )
    last_time_value: float | None = None
    lock: Lock = field(default_factory=Lock)

    def reset_stream(self) -> None:
        self.data = {"time": [], "x": [], "y": [], "z": [], "mag": []}
        self.last_time_value = None
        self.session_id = None


STATE = StreamState()


def normalize_base_url(raw_url: str) -> str:
    raw_url = raw_url.strip()
    if not raw_url:
        raise ValueError("Bitte eine phyphox-URL angeben.")
    if not raw_url.startswith(("http://", "https://")):
        raw_url = f"http://{raw_url}"
    parsed = urlparse(raw_url)
    if not parsed.scheme or not parsed.netloc:
        raise ValueError("Die phyphox-URL ist ungültig.")
    return raw_url.rstrip("/")


def request_json(base_url: str, path: str, params: dict[str, Any] | None = None) -> dict[str, Any]:
    response = requests.get(f"{base_url}{path}", params=params, timeout=4)
    response.raise_for_status()
    return response.json()


def request_control(base_url: str, command: str) -> dict[str, Any]:
    response = requests.get(
        f"{base_url}/control", params={"cmd": command}, timeout=4
    )
    response.raise_for_status()
    return response.json()


TIME_PATTERNS = [
    re.compile(r"^t$"),
    re.compile(r"^time$"),
    re.compile(r"time"),
    re.compile(r"timestamp"),
]
X_PATTERNS = [
    re.compile(r"^accx$"),
    re.compile(r"^x$"),
    re.compile(r"acc.*x"),
    re.compile(r"x.*acc"),
    re.compile(r"lin.*x"),
]
Y_PATTERNS = [
    re.compile(r"^accy$"),
    re.compile(r"^y$"),
    re.compile(r"acc.*y"),
    re.compile(r"y.*acc"),
    re.compile(r"lin.*y"),
]
Z_PATTERNS = [
    re.compile(r"^accz$"),
    re.compile(r"^z$"),
    re.compile(r"acc.*z"),
    re.compile(r"z.*acc"),
    re.compile(r"lin.*z"),
]


def score_name(name: str, patterns: list[re.Pattern[str]]) -> int:
    cleaned = re.sub(r"[^a-z0-9]", "", name.lower())
    score = 0
    for index, pattern in enumerate(patterns):
        if pattern.search(cleaned):
            score = max(score, 100 - index * 10)
    return score


def auto_detect_buffers(buffer_names: list[str]) -> dict[str, str | None]:
    if not buffer_names:
        return {"time": None, "x": None, "y": None, "z": None}

    def best(patterns: list[re.Pattern[str]], excluded: set[str] | None = None) -> str | None:
        excluded = excluded or set()
        ranked = sorted(
            ((score_name(name, patterns), name) for name in buffer_names if name not in excluded),
            reverse=True,
        )
        if not ranked or ranked[0][0] <= 0:
            return None
        return ranked[0][1]

    detected: dict[str, str | None] = {"time": None, "x": None, "y": None, "z": None}
    used: set[str] = set()

    detected["time"] = best(TIME_PATTERNS, used)
    if detected["time"]:
        used.add(detected["time"])
    detected["x"] = best(X_PATTERNS, used)
    if detected["x"]:
        used.add(detected["x"])
    detected["y"] = best(Y_PATTERNS, used)
    if detected["y"]:
        used.add(detected["y"])
    detected["z"] = best(Z_PATTERNS, used)
    return detected


def current_buffer_map() -> dict[str, str]:
    missing = [k for k, v in STATE.buffers.items() if not v]
    if missing:
        raise ValueError("Zeit- und Beschleunigungsbuffer müssen gesetzt sein.")
    return {k: str(v) for k, v in STATE.buffers.items()}


def build_poll_params() -> dict[str, str]:
    buffers = current_buffer_map()
    time_name = buffers["time"]
    x_name = buffers["x"]
    y_name = buffers["y"]
    z_name = buffers["z"]

    if STATE.last_time_value is None:
        return {
            time_name: "full",
            x_name: "full",
            y_name: "full",
            z_name: "full",
        }

    threshold = repr(STATE.last_time_value)
    return {
        time_name: threshold,
        x_name: f"{threshold}|{time_name}",
        y_name: f"{threshold}|{time_name}",
        z_name: f"{threshold}|{time_name}",
    }


def to_float_list(values: list[Any]) -> list[float]:
    converted: list[float] = []
    for value in values:
        try:
            converted.append(float(value))
        except (TypeError, ValueError):
            continue
    return converted


def align_series(values: list[float], length: int) -> list[float | None]:
    if len(values) >= length:
        return values[:length]
    return values + [None] * (length - len(values))


def append_chunk(time_values: list[float], x_values: list[float], y_values: list[float], z_values: list[float]) -> dict[str, Any]:
    incoming = list(zip(time_values, align_series(x_values, len(time_values)), align_series(y_values, len(time_values)), align_series(z_values, len(time_values))))
    if STATE.last_time_value is not None:
        incoming = [row for row in incoming if row[0] > STATE.last_time_value]

    if not incoming:
        latest = None
        if STATE.data["time"]:
            latest = {
                "time": STATE.data["time"][-1],
                "x": STATE.data["x"][-1],
                "y": STATE.data["y"][-1],
                "z": STATE.data["z"][-1],
                "mag": STATE.data["mag"][-1],
            }
        return {
            "chunk": {"time": [], "x": [], "y": [], "z": [], "mag": []},
            "count": len(STATE.data["time"]),
            "latest": latest,
        }

    time_values = [row[0] for row in incoming]
    x_aligned = [row[1] for row in incoming]
    y_aligned = [row[2] for row in incoming]
    z_aligned = [row[3] for row in incoming]
    n = len(time_values)
    if n == 0:
        return {
            "chunk": {"time": [], "x": [], "y": [], "z": [], "mag": []},
            "count": len(STATE.data["time"]),
            "latest": None,
        }

    magnitude: list[float | None] = []
    for x, y, z in zip(x_aligned, y_aligned, z_aligned):
        if x is None or y is None or z is None:
            magnitude.append(None)
        else:
            magnitude.append(math.sqrt(x * x + y * y + z * z))

    STATE.data["time"].extend(time_values)
    STATE.data["x"].extend(x_aligned)
    STATE.data["y"].extend(y_aligned)
    STATE.data["z"].extend(z_aligned)
    STATE.data["mag"].extend(magnitude)
    STATE.last_time_value = time_values[-1]

    latest = {
        "time": STATE.data["time"][-1],
        "x": STATE.data["x"][-1],
        "y": STATE.data["y"][-1],
        "z": STATE.data["z"][-1],
        "mag": STATE.data["mag"][-1],
    }

    return {
        "chunk": {
            "time": time_values,
            "x": x_aligned,
            "y": y_aligned,
            "z": z_aligned,
            "mag": magnitude,
        },
        "count": len(STATE.data["time"]),
        "latest": latest,
    }


def build_metadata_sheet(workbook: Workbook, start_time: float | None, end_time: float | None, row_count: int) -> None:
    ws = workbook.create_sheet("Metadaten")
    ws["A1"] = "Schlüssel"
    ws["B1"] = "Wert"
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)

    items = [
        ("Exportzeitpunkt", datetime.now().isoformat(timespec="seconds")),
        ("phyphox-URL", STATE.base_url or ""),
        ("Experiment", STATE.experiment_title or ""),
        ("Zeit-Buffer", STATE.buffers.get("time") or ""),
        ("X-Buffer", STATE.buffers.get("x") or ""),
        ("Y-Buffer", STATE.buffers.get("y") or ""),
        ("Z-Buffer", STATE.buffers.get("z") or ""),
        ("Ausgewählter Start", "" if start_time is None else start_time),
        ("Ausgewähltes Ende", "" if end_time is None else end_time),
        ("Exportierte Zeilen", row_count),
    ]

    for idx, (key, value) in enumerate(items, start=2):
        ws[f"A{idx}"] = key
        ws[f"B{idx}"] = value

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 28


@app.route("/")
def index() -> str:
    return render_template("index.html")


@app.route("/api/connect", methods=["POST"])
def api_connect():
    payload = request.get_json(silent=True) or {}
    base_url_raw = str(payload.get("base_url", ""))

    try:
        base_url = normalize_base_url(base_url_raw)
        config = request_json(base_url, "/config")
        buffer_entries = config.get("buffers", [])
        buffer_names = [entry.get("name") for entry in buffer_entries if isinstance(entry, dict) and entry.get("name")]
        detected = auto_detect_buffers(buffer_names)

        with STATE.lock:
            STATE.base_url = base_url
            STATE.experiment_title = config.get("localTitle") or config.get("title") or "phyphox"
            STATE.buffer_options = buffer_names
            STATE.buffers.update(detected)
            STATE.connected = True
            STATE.reset_stream()

        return jsonify(
            {
                "ok": True,
                "base_url": STATE.base_url,
                "experiment_title": STATE.experiment_title,
                "buffer_options": STATE.buffer_options,
                "detected": STATE.buffers,
            }
        )
    except requests.RequestException as exc:
        return jsonify({"ok": False, "error": f"Verbindung zu phyphox fehlgeschlagen: {exc}"}), 400
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.route("/api/buffers", methods=["POST"])
def api_set_buffers():
    payload = request.get_json(silent=True) or {}
    with STATE.lock:
        if not STATE.connected:
            return jsonify({"ok": False, "error": "Es besteht noch keine Verbindung zu phyphox."}), 400

        for key in ("time", "x", "y", "z"):
            value = payload.get(key)
            STATE.buffers[key] = str(value).strip() if value else None
        STATE.reset_stream()

    missing = [k for k, v in STATE.buffers.items() if not v]
    if missing:
        return jsonify({"ok": False, "error": "Bitte alle vier Buffer auswählen."}), 400

    return jsonify({"ok": True, "buffers": STATE.buffers})


@app.route("/api/control/<command>", methods=["POST"])
def api_control(command: str):
    if command not in {"start", "stop", "clear"}:
        return jsonify({"ok": False, "error": "Unbekannter Befehl."}), 400

    with STATE.lock:
        if not STATE.base_url:
            return jsonify({"ok": False, "error": "Keine phyphox-Verbindung vorhanden."}), 400
        try:
            result = request_control(STATE.base_url, command)
            if command == "clear":
                STATE.reset_stream()
            return jsonify({"ok": bool(result.get("result", False))})
        except requests.RequestException as exc:
            return jsonify({"ok": False, "error": f"Steuerbefehl fehlgeschlagen: {exc}"}), 400


@app.route("/api/poll")
def api_poll():
    with STATE.lock:
        if not STATE.connected or not STATE.base_url:
            return jsonify({"ok": False, "error": "Keine phyphox-Verbindung vorhanden."}), 400
        try:
            params = build_poll_params()
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        base_url = STATE.base_url
        buffers = current_buffer_map()

    try:
        response = request_json(base_url, "/get", params=params)
    except requests.RequestException as exc:
        return jsonify({"ok": False, "error": f"Abruf fehlgeschlagen: {exc}"}), 400

    with STATE.lock:
        status = response.get("status", {})
        session = status.get("session")
        if STATE.session_id and session and session != STATE.session_id:
            STATE.reset_stream()
        if session:
            STATE.session_id = session

        buffer_block = response.get("buffer", {})
        time_values = to_float_list(buffer_block.get(buffers["time"], {}).get("buffer", []))
        x_values = to_float_list(buffer_block.get(buffers["x"], {}).get("buffer", []))
        y_values = to_float_list(buffer_block.get(buffers["y"], {}).get("buffer", []))
        z_values = to_float_list(buffer_block.get(buffers["z"], {}).get("buffer", []))

        update = append_chunk(time_values, x_values, y_values, z_values)
        update["ok"] = True
        update["status"] = {
            "measuring": bool(status.get("measuring", False)),
            "timedRun": bool(status.get("timedRun", False)),
            "countDown": status.get("countDown"),
            "session": STATE.session_id,
        }
        return jsonify(update)


@app.route("/api/export", methods=["POST"])
def api_export():
    payload = request.get_json(silent=True) or {}
    start_time = payload.get("start")
    end_time = payload.get("end")

    try:
        start_value = None if start_time in (None, "", "null") else float(start_time)
        end_value = None if end_time in (None, "", "null") else float(end_time)
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Ungültiger Zeitbereich."}), 400

    with STATE.lock:
        time_values = STATE.data["time"]
        if not time_values:
            return jsonify({"ok": False, "error": "Es liegen noch keine Messdaten vor."}), 400

        rows: list[tuple[float, float | None, float | None, float | None, float | None]] = []
        for t, x, y, z, mag in zip(
            STATE.data["time"],
            STATE.data["x"],
            STATE.data["y"],
            STATE.data["z"],
            STATE.data["mag"],
        ):
            if start_value is not None and t < start_value:
                continue
            if end_value is not None and t > end_value:
                continue
            rows.append((t, x, y, z, mag))

    if not rows:
        return jsonify({"ok": False, "error": "Im gewählten Zeitbereich liegen keine Daten."}), 400

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Messdaten"
    headers = ["time_s", "acc_x", "acc_y", "acc_z", "acc_abs"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col_index, value=value)

    for column in ["A", "B", "C", "D", "E"]:
        ws.column_dimensions[column].width = 14

    build_metadata_sheet(workbook, start_value, end_value, len(rows))

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp_path = Path(tmp.name)
    tmp.close()
    workbook.save(tmp_path)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"phyphox_export_{timestamp}.xlsx"
    return send_file(
        tmp_path,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/state")
def api_state():
    with STATE.lock:
        return jsonify(
            {
                "ok": True,
                "connected": STATE.connected,
                "base_url": STATE.base_url,
                "experiment_title": STATE.experiment_title,
                "buffer_options": STATE.buffer_options,
                "buffers": STATE.buffers,
                "count": len(STATE.data["time"]),
            }
        )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
